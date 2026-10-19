"""Reads the vehicle catalogue and the occupations table from the challenge workbook.

Use value, drive and list price are deliberately not read from the workbook:
they come from the Tax Authority CSV through the Step 1 lookup (lookup_vehicle).
"""

from collections.abc import Callable, Mapping
from dataclasses import dataclass
from functools import cache
from pathlib import Path
from types import MappingProxyType
from typing import Any, TypeVar

import openpyxl

from app.config_file import shown
from app.errors import ConfigError
from app.lookup_config import LookupConfig
from app.models import VehicleRecord
from app.vehicle_repository import VehicleRepository
from app.workbook_columns import (
    CATALOG_COLUMN_TO_FIELD,
    CATALOG_SHEET,
    NO,
    OCCUPATION_COLUMN_TO_FIELD,
    OCCUPATIONS_SHEET,
    YES,
)

DATA_DIR = Path(__file__).parent.parent / "data"

T = TypeVar("T")


@dataclass(frozen=True, slots=True)
class CatalogCar:
    car_id: str
    group: str
    manufacturer_code: int
    commercial_name: str
    model_code: int
    registration_year: int
    monthly_cost: int
    default_participation: int


@dataclass(frozen=True, slots=True)
class Occupation:
    name: str
    group: str
    car_required: bool
    waiver_allowed: bool
    waiver_maintenance: int  # 0 when waiving the car is not allowed


@dataclass(frozen=True, slots=True)
class Catalog:
    cars: Mapping[str, CatalogCar]  # by car_id, in workbook order
    occupations: Mapping[str, Occupation]  # by name, in workbook order


@cache
def find_workbook(data_dir: Path = DATA_DIR) -> Path:
    """The single .xlsx in data_dir that has the catalogue sheet."""
    matches = [p for p in sorted(data_dir.glob("*.xlsx")) if CATALOG_SHEET in _sheet_names(p)]
    if len(matches) != 1:
        raise ConfigError(
            "*.xlsx", f"exactly one workbook with sheet '{CATALOG_SHEET}'", shown([p.name for p in matches]), file=data_dir
        )
    return matches[0]


def load_catalog(path: Path | None = None) -> Catalog:
    """Read catalogue and occupations; later calls with the same workbook return the same object."""
    return _load((path or find_workbook()).resolve())


def lookup_vehicle(car: CatalogCar, repo: VehicleRepository, lookup: LookupConfig) -> VehicleRecord:
    """Find the catalogue car in the use-value CSV by the official key (Step 1)."""
    known = {
        "tax_year": lookup.tax_year,
        "registration_year": car.registration_year,
        "vehicle_type_code": lookup.catalogue_vehicle_type,
        "manufacturer_code": car.manufacturer_code,
        "model_code": car.model_code,
    }
    unknown = [field for field in repo.key_fields if field not in known]
    if unknown:
        raise ValueError(f"Lookup key fields {unknown} have no value for catalogue car {car.car_id}")
    return repo.find_exact(**{field: known[field] for field in repo.key_fields})


@cache
def _load(path: Path) -> Catalog:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        cars = _read_sheet(workbook, CATALOG_SHEET, CATALOG_COLUMN_TO_FIELD, _to_car)
        occupations = _read_sheet(workbook, OCCUPATIONS_SHEET, OCCUPATION_COLUMN_TO_FIELD, _to_occupation)
        return Catalog(
            cars=_unique_by(cars, "car_id", CATALOG_SHEET),
            occupations=_unique_by(occupations, "name", OCCUPATIONS_SHEET),
        )
    except ConfigError as e:
        e.file = path
        raise
    finally:
        workbook.close()


def _sheet_names(path: Path) -> list[str]:
    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


def _read_sheet(workbook: Any, sheet: str, column_to_field: dict[str, str], build: Callable[["_Row"], T]) -> list[T]:
    if sheet not in workbook.sheetnames:
        raise ConfigError(sheet, "a sheet with this name", f"sheets {shown(workbook.sheetnames)}")
    header, *rows = workbook[sheet].iter_rows(values_only=True)
    columns = [str(name).strip() if name is not None else "" for name in header]

    missing = [column for column in column_to_field if column not in columns]
    if missing:
        raise ConfigError(sheet, f"columns {shown(list(column_to_field))}", f"missing {shown(missing)}")

    index = {column_to_field[c]: columns.index(c) for c in column_to_field}
    return [
        build(_Row(sheet, number, {field: row[i] for field, i in index.items()}, column_to_field))
        for number, row in enumerate(rows, start=2)  # row 1 is the header
        if row and row[0] is not None
    ]


def _to_car(row: "_Row") -> CatalogCar:
    return CatalogCar(
        car_id=row.text("car_id"),
        group=row.text("group"),
        manufacturer_code=row.code("manufacturer_code"),
        commercial_name=row.text("commercial_name"),
        model_code=row.code("model_code"),
        registration_year=row.amount("registration_year", minimum=1),
        monthly_cost=row.amount("monthly_cost", minimum=1),
        default_participation=row.amount("default_participation", minimum=0),
    )


def _to_occupation(row: "_Row") -> Occupation:
    waiver_allowed = row.yes_no("waiver_allowed")
    maintenance = row.amount("waiver_maintenance", minimum=0)
    if waiver_allowed and maintenance == 0:
        raise row.error("waiver_maintenance", "an amount > 0 when waiving is allowed", maintenance)
    if not waiver_allowed and maintenance != 0:
        raise row.error("waiver_maintenance", "0 when waiving is not allowed", maintenance)
    return Occupation(
        name=row.text("name"),
        group=row.text("group"),
        car_required=row.yes_no("car_required"),
        waiver_allowed=waiver_allowed,
        waiver_maintenance=maintenance,
    )


def _unique_by(items: list[T], field: str, sheet: str) -> Mapping[str, T]:
    by_key: dict[str, T] = {}
    for item in items:
        key = getattr(item, field)
        if key in by_key:
            raise ConfigError(sheet, f"a unique value in every row for {field}", f"{shown(key)} twice")
        by_key[key] = item
    return MappingProxyType(by_key)


class _Row:
    """One workbook row plus where it came from, so errors can name sheet, row and column."""

    def __init__(self, sheet: str, number: int, values: dict[str, Any], column_to_field: dict[str, str]) -> None:
        self._sheet = sheet
        self._number = number
        self._values = values
        self._column_of = {field: column for column, field in column_to_field.items()}

    def text(self, field: str) -> str:
        value = self._values[field]
        if not isinstance(value, str) or not value.strip():
            raise self.error(field, "non-empty text", value)
        return value.strip()

    def code(self, field: str) -> int:
        """Codes are stored as zero-padded text ('00885') or as numbers."""
        value = self._values[field]
        if isinstance(value, int) and not isinstance(value, bool):
            return value
        if isinstance(value, str) and value.strip().isdigit():
            return int(value)
        raise self.error(field, "a numeric code", value)

    def amount(self, field: str, *, minimum: int) -> int:
        value = self._values[field]
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        if not isinstance(value, int) or isinstance(value, bool) or value < minimum:
            raise self.error(field, f"a whole number >= {minimum}", value)
        return value

    def yes_no(self, field: str) -> bool:
        value = self._values[field]
        if value not in (YES, NO):
            raise self.error(field, f"'{YES}' or '{NO}'", value)
        return value == YES

    def error(self, field: str, expected: str, found: Any) -> ConfigError:
        return ConfigError(f"{self._sheet}[row {self._number}].{self._column_of[field]}", expected, shown(found))
