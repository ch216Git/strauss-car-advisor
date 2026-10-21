"""Details from a file the employee uploads (payslip, form, table, photo), kept only when certain.

1. Python reads the file's text (PDF text layer, TXT/CSV, Excel). Images and scanned PDFs have no
   text to read; then the file itself is given to the reader (a multimodal model).
2. The reader (app.agent.providers.gemini.GeminiDocumentReader) returns each field it found with the
   exact quote it read it from, and the fields it found unclear.
3. Python decides what is certain. A value is kept only when: it parses and is in range; it is written
   in its quote; the quote is in the file's text (when the file has text); the field has one value;
   an occupation matches exactly one workbook occupation; a car identifies exactly one catalogue car.
   Everything else is reported as unclear. Nothing is calculated, converted or guessed.
4. Nothing is saved here. The employee reviews what was found and confirms it in the chat, in their
   own words (frontend), so the advisor saves it with the usual tools and grounding checks.

The file is never stored and its content is never logged.
"""

import csv
import io
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from typing import Any, Protocol

from app.agent.grounding import matching_occupations, number_grounded
from app.agent.tools import ToolContext
from app.agent.views import car_view
from app.car_search import MatchStatus, normalize, prefix_stripped

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 30_000  # what is sent to the reader; a payslip or form is far shorter
MIN_TEXT_CHARS = 20  # a PDF with less text than this is treated as a scan

# What may be read, in presentation order, with the Hebrew label the employee sees.
FIELDS: dict[str, str] = {
    "occupation": "תפקיד",
    "base_salary": "שכר בסיס חודשי (ברוטו)",
    "commissions": "עמלות חודשיות",
    "bonus": "בונוס חודשי",
    "pension_percent": "הפרשה לפנסיה",
    "training_fund_percent": "הפרשה לקרן השתלמות",
    "credit_points": "נקודות זיכוי",
    "eligible_children": "מספר ילדים",
    "current_car": "הרכב היום",
    "current_participation": "השתתפות עצמית חודשית היום",
}
_SHEKELS = {"base_salary", "commissions", "bonus", "current_participation"}
_PERCENT = {"pension_percent", "training_fund_percent"}
# (minimum, maximum) a value may have; outside it the reading is not trusted.
_RANGES: dict[str, tuple[float, float]] = {
    "base_salary": (1, 1_000_000), "commissions": (0, 1_000_000), "bonus": (0, 1_000_000),
    "pension_percent": (0, 30), "training_fund_percent": (0, 30), "credit_points": (0, 20),
    "eligible_children": (0, 20), "current_participation": (0, 100_000),
}

# File kinds by extension: (mime type, how its text is read).
KINDS: dict[str, tuple[str, str]] = {
    "pdf": ("application/pdf", "pdf"),
    "png": ("image/png", "image"),
    "jpg": ("image/jpeg", "image"),
    "jpeg": ("image/jpeg", "image"),
    "webp": ("image/webp", "image"),
    "txt": ("text/plain", "text"),
    "csv": ("text/csv", "csv"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
}
# The first bytes each binary kind must start with: the extension alone is not trusted.
_MAGIC = {
    "pdf": (b"%PDF",), "png": (b"\x89PNG",), "jpg": (b"\xff\xd8\xff",), "jpeg": (b"\xff\xd8\xff",),
    "webp": (b"RIFF",), "xlsx": (b"PK\x03\x04",),
}


class DocumentError(ValueError):
    """The file cannot be used. code is a stable API error code."""

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


@dataclass(frozen=True, slots=True)
class Document:
    name: str
    data: bytes

    @property
    def extension(self) -> str:
        return self.name.rsplit(".", 1)[-1].lower() if "." in self.name else ""

    @property
    def mime_type(self) -> str:
        return KINDS[self.extension][0]


class DocumentReader(Protocol):
    def read(self, text: str | None, document: Document) -> Mapping[str, Any]:
        """{"fields": [{"field", "value", "quote"}], "unclear": [field]} for text, or (text None) the file."""
        ...


@dataclass(frozen=True, slots=True)
class Found:
    field: str
    label: str
    value: str  # for display and for the confirmation message, e.g. "18,500 ₪", "6%", "מנהל/ת מוצר"


@dataclass(frozen=True, slots=True)
class DocumentDetails:
    found: tuple[Found, ...]
    unclear: tuple[str, ...]  # field names, in FIELDS order
    from_image: bool  # read from an image or a scan: the quotes could not be checked against text


def check_document(document: Document) -> None:
    """Raises DocumentError for an empty, oversized, unsupported or mislabelled file."""
    if not document.data:
        raise DocumentError("unreadable_file")
    if len(document.data) > MAX_FILE_BYTES:
        raise DocumentError("file_too_large")
    if document.extension not in KINDS:
        raise DocumentError("unsupported_file")
    magic = _MAGIC.get(document.extension)
    if magic and not document.data.startswith(magic):
        raise DocumentError("unsupported_file")


def document_text(document: Document) -> str | None:
    """The file's text, or None when it has none (an image or a scanned PDF). Raises DocumentError."""
    kind = KINDS[document.extension][1]
    try:
        if kind == "image":
            return None
        if kind == "text":
            text = _decode(document.data)
        elif kind == "csv":
            rows = csv.reader(io.StringIO(_decode(document.data)))
            text = "\n".join(" | ".join(cell.strip() for cell in row if cell.strip()) for row in rows)
        elif kind == "xlsx":
            text = _xlsx_text(document.data)
        else:
            text = _pdf_text(document.data)
            if len(text.strip()) < MIN_TEXT_CHARS:
                return None  # a scan: read as an image
    except DocumentError:
        raise
    except Exception:  # a corrupt file; its content is never echoed
        raise DocumentError("unreadable_file") from None
    if not text.strip():
        raise DocumentError("unreadable_file")
    return text[:MAX_TEXT_CHARS]


def read_details(document: Document, reader: DocumentReader, ctx: ToolContext) -> DocumentDetails:
    check_document(document)
    text = document_text(document)
    raw = reader.read(text, document)
    return verified(raw, text, ctx)


def verified(raw: Mapping[str, Any], text: str | None, ctx: ToolContext) -> DocumentDetails:
    """Only what Python can confirm; every other field the reader mentioned is unclear."""
    entries = raw.get("fields") if isinstance(raw, Mapping) else None
    unclear = {f for f in (raw.get("unclear") or ()) if f in FIELDS} if isinstance(raw, Mapping) else set()
    values: dict[str, set[str]] = {}
    for entry in entries if isinstance(entries, Sequence) else ():
        if not isinstance(entry, Mapping) or entry.get("field") not in FIELDS:
            continue
        field = entry["field"]
        value = _checked(field, str(entry.get("value", "")), str(entry.get("quote", "")), text, ctx)
        if value is None:
            unclear.add(field)
        else:
            values.setdefault(field, set()).add(value)
    found = []
    for field, label in FIELDS.items():
        readings = values.get(field, set())
        if len(readings) == 1 and field not in unclear:
            found.append(Found(field, label, readings.pop()))
        elif readings:  # two different values for one field, or the reader also called it unclear
            unclear.add(field)
    return DocumentDetails(tuple(found), tuple(f for f in FIELDS if f in unclear), from_image=text is None)


# --- Checks ---------------------------------------------------------------------------------------

def _checked(field: str, value: str, quote: str, text: str | None, ctx: ToolContext) -> str | None:
    """The value as shown to the employee, or None if it is not certain."""
    value, quote = value.strip(), quote.strip()
    if not value or not quote or len(quote) > 300:
        return None
    if text is not None and _squeezed(quote) not in _squeezed(text):
        return None  # the quote is not in the file
    if field == "occupation":
        names = matching_occupations(value, ctx.catalog.occupations)
        if len(names) != 1 or not _words_in(value, quote):
            return None
        return names[0]
    if field == "current_car":
        found = ctx.car_search.search(value)
        if found.status != MatchStatus.IDENTIFIED or len(found.car_ids) != 1 or not _words_in(value, quote):
            return None
        view = car_view(found.car_ids[0], ctx)
        return f"{view.manufacturer} {view.commercial_name} ({view.registration_year})"
    number = _number(value)
    if number is None or not number_grounded(number, (quote,)):
        return None
    low, high = _RANGES[field]
    if not low <= number <= high:
        return None
    if field == "eligible_children" and number != int(number):
        return None
    return _display(field, number)


def _number(value: str) -> float | None:
    cleaned = re.sub(r"[₪%\s]|ש\"ח|ש״ח|שח", "", value).replace(",", "")
    try:
        return float(cleaned)
    except ValueError:
        return None


def _display(field: str, number: float) -> str:
    shown = f"{int(number):,}" if number == int(number) else f"{number:,}".rstrip("0").rstrip(".")
    if field in _SHEKELS:
        return f"{shown} ₪"
    if field in _PERCENT:
        return f"{shown}%"
    return shown


def _words_in(value: str, quote: str) -> bool:
    """Every word of value is written in quote (a Hebrew prefix letter allowed: "בקורולה")."""
    quoted = {form for token in normalize(quote) for form in (token, *prefix_stripped(token))}
    return all(word in quoted for word in normalize(value))


_INVISIBLE = re.compile(r"[‎‏‪-‮⁦-⁩]")


def _squeezed(text: str) -> str:
    """Whitespace and direction marks ignored, so a quote matches however the file's text was laid out."""
    return re.sub(r"\s+", "", _INVISIBLE.sub("", text)).casefold()


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1255"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentError("unreadable_file")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise DocumentError("unreadable_file")
    return "\n".join(page.extract_text() or "" for page in reader.pages[:10])


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for sheet in workbook.worksheets[:5]:
        for row in sheet.iter_rows(max_row=500, values_only=True):
            cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if cells:
                lines.append(" | ".join(cells))
    workbook.close()
    return "\n".join(lines)
