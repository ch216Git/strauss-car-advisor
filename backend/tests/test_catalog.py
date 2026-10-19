from dataclasses import FrozenInstanceError, fields
from pathlib import Path

import openpyxl
import pytest

from app.catalog import Catalog, CatalogCar, find_workbook, load_catalog, lookup_vehicle
from app.errors import ConfigError
from app.lookup_config import load_lookup_config
from app.rules import load_rules
from app.vehicle_repository import VehicleRepository
from app.workbook_columns import CATALOG_CROSS_CHECK_COLUMNS, CATALOG_SHEET, OCCUPATIONS_SHEET

# Hand-copied from the challenge brief and workbook ('כללים ונתונים', 'קטלוג רכבים', 'עיסוקים');
# an independent oracle for the reading, not business logic.
BRIEF_CAR_COUNT = 20  # "20 רכבים ... + 2 במסלול המשפחתי": the sheet has CAR-01..CAR-20, family cars included
BRIEF_OCCUPATION_COUNT = 14
FAMILY_CAR_COSTS = {"CAR-19": 6000, "CAR-20": 5600}  # קרניבל 6,000 ₪, סנטה פה 5,600 ₪


@pytest.fixture(scope="module")
def catalog() -> Catalog:
    return load_catalog()


@pytest.fixture(scope="module")
def repo() -> VehicleRepository:
    return VehicleRepository()


# --- Loading -----------------------------------------------------------------

def test_catalog_loads(catalog: Catalog) -> None:
    assert len(catalog.cars) == BRIEF_CAR_COUNT
    assert len(catalog.occupations) == BRIEF_OCCUPATION_COUNT
    assert load_catalog() is catalog


def test_car_fields(catalog: Catalog) -> None:
    carnival = catalog.cars["CAR-19"]

    assert (carnival.commercial_name, carnival.manufacturer_code, carnival.model_code, carnival.registration_year) == (
        "CARNIVAL", 885, 648, 2025
    )
    assert {car_id: catalog.cars[car_id].monthly_cost for car_id in FAMILY_CAR_COSTS} == FAMILY_CAR_COSTS
    assert carnival.default_participation == 500


def test_occupation_fields(catalog: Catalog) -> None:
    field_rep = catalog.occupations["נציג/ת מכירות שטח"]
    developer = catalog.occupations["מפתח/ת תוכנה"]

    assert (field_rep.group, field_rep.car_required, field_rep.waiver_allowed, field_rep.waiver_maintenance) == ("ב", True, False, 0)
    assert (developer.group, developer.car_required, developer.waiver_allowed, developer.waiver_maintenance) == ("ב", False, True, 1500)


def test_catalog_is_immutable(catalog: Catalog) -> None:
    with pytest.raises(FrozenInstanceError):
        catalog.cars["CAR-01"].monthly_cost = 0  # type: ignore[misc]
    with pytest.raises(TypeError):
        catalog.cars["CAR-99"] = catalog.cars["CAR-01"]  # type: ignore[index]


def test_catalog_car_has_no_values_owned_by_the_csv() -> None:
    # Use value, drive and list price have one source: the Tax Authority CSV via lookup_vehicle.
    assert not {f.name for f in fields(CatalogCar)} & set(CATALOG_CROSS_CHECK_COLUMNS.values())


# --- Consistency with rules.json ---------------------------------------------

def test_family_track_cars_exist_in_catalog(catalog: Catalog) -> None:
    assert set(load_rules().family_track.cars) <= set(catalog.cars)


def test_every_regular_car_and_occupation_has_a_rules_group(catalog: Catalog) -> None:
    rules = load_rules()
    regular_cars = [car for car in catalog.cars.values() if car.car_id not in rules.family_track.cars]

    assert {car.group for car in regular_cars} <= set(rules.group_order)
    assert {occupation.group for occupation in catalog.occupations.values()} <= set(rules.group_order)


# --- Use value through the Step 1 lookup -------------------------------------

def test_every_catalog_car_resolves_in_the_csv(catalog: Catalog, repo: VehicleRepository) -> None:
    lookup = load_lookup_config()

    for car in catalog.cars.values():
        record = lookup_vehicle(car, repo, lookup)
        assert (record.manufacturer_code, record.model_code, record.registration_year) == (
            car.manufacturer_code, car.model_code, car.registration_year
        )
        assert record.tax_year == lookup.tax_year
        assert record.vehicle_type_code == lookup.catalogue_vehicle_type


def test_vehicle_type_assumption_resolves_atto_3(catalog: Catalog, repo: VehicleRepository) -> None:
    # Under vehicle type 2 the same key is the BYD TANG; the configured type 1 must give the ATTO 3.
    record = lookup_vehicle(catalog.cars["CAR-10"], repo, load_lookup_config())

    assert record.commercial_name == "ATTO 3"


def test_use_value_comes_from_csv_and_agrees_with_workbook(catalog: Catalog, repo: VehicleRepository) -> None:
    # The workbook column is read here only as a test oracle (see CATALOG_CROSS_CHECK_COLUMNS).
    columns = {"מזהה": "car_id", "שווי שימוש חודשי": "use_value"}
    workbook_use_values = {row["car_id"]: row["use_value"] for row in read_rows(find_workbook(), CATALOG_SHEET, columns)}
    lookup = load_lookup_config()

    csv_use_values = {car.car_id: lookup_vehicle(car, repo, lookup).use_value for car in catalog.cars.values()}

    assert csv_use_values == workbook_use_values


# --- Workbook errors ---------------------------------------------------------

def read_rows(path: Path, sheet: str, columns: dict[str, str]) -> list[dict]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    header, *rows = workbook[sheet].iter_rows(values_only=True)
    workbook.close()
    index = {field: header.index(column) for column, field in columns.items()}
    return [{field: row[i] for field, i in index.items()} for row in rows if row[0]]


def copy_workbook(tmp_path: Path, change=None) -> Path:
    """Copy both sheets of the real workbook into a new file, letting `change` edit the rows first."""
    source = openpyxl.load_workbook(find_workbook(), read_only=True, data_only=True)
    sheets = {name: [list(row) for row in source[name].iter_rows(values_only=True)] for name in (CATALOG_SHEET, OCCUPATIONS_SHEET)}
    source.close()
    if change:
        change(sheets)
    target = openpyxl.Workbook()
    target.remove(target.active)
    for name, rows in sheets.items():
        sheet = target.create_sheet(name)
        for row in rows:
            sheet.append(row)
    path = tmp_path / "workbook.xlsx"
    target.save(path)
    return path


def set_cell(sheet: str, row: int, column: str, value):
    def change(sheets):
        header = sheets[sheet][0]
        sheets[sheet][row - 1][header.index(column)] = value
    return change


def catalog_error(path: Path) -> ConfigError:
    with pytest.raises(ConfigError) as info:
        load_catalog(path)
    assert info.value.file == path.resolve()
    return info.value


def test_copied_workbook_loads(tmp_path: Path) -> None:
    assert load_catalog(copy_workbook(tmp_path)).cars == load_catalog().cars


@pytest.mark.parametrize(
    "change, path, expected, found",
    [
        (set_cell(CATALOG_SHEET, 3, "קוד תוצר", "abc"), "קטלוג רכבים[row 3].קוד תוצר", "a numeric code", '"abc"'),
        (set_cell(CATALOG_SHEET, 2, "עלות חודשית לחברה (דמה)", 0), "קטלוג רכבים[row 2].עלות חודשית לחברה (דמה)", "a whole number >= 1", "0"),
        (set_cell(CATALOG_SHEET, 2, "קבוצה", None), "קטלוג רכבים[row 2].קבוצה", "non-empty text", "null"),
        (set_cell(CATALOG_SHEET, 3, "מזהה", "CAR-01"), CATALOG_SHEET, "a unique value in every row for car_id", '"CAR-01" twice'),
        (set_cell(OCCUPATIONS_SHEET, 2, "רכב חובה", "yes"), "עיסוקים[row 2].רכב חובה", "'כן' or 'לא'", '"yes"'),
        (set_cell(OCCUPATIONS_SHEET, 6, "אחזקת רכב בוויתור (₪)", 0), "עיסוקים[row 6].אחזקת רכב בוויתור (₪)", "an amount > 0 when waiving is allowed", "0"),
        (set_cell(OCCUPATIONS_SHEET, 2, "אחזקת רכב בוויתור (₪)", 900), "עיסוקים[row 2].אחזקת רכב בוויתור (₪)", "0 when waiving is not allowed", "900"),
    ],
)
def test_invalid_workbook_value_raises(tmp_path: Path, change, path: str, expected: str, found: str) -> None:
    error = catalog_error(copy_workbook(tmp_path, change))

    assert (error.path, error.expected, error.found) == (path, expected, found)


def test_missing_column_raises(tmp_path: Path) -> None:
    error = catalog_error(copy_workbook(tmp_path, set_cell(CATALOG_SHEET, 1, "קוד דגם", "דגם")))

    assert (error.path, error.found) == (CATALOG_SHEET, 'missing ["קוד דגם"]')


def test_missing_sheet_raises(tmp_path: Path) -> None:
    error = catalog_error(copy_workbook(tmp_path, lambda sheets: sheets.pop(OCCUPATIONS_SHEET)))

    assert (error.path, error.expected) == (OCCUPATIONS_SHEET, "a sheet with this name")


def test_no_workbook_in_data_dir_raises(tmp_path: Path) -> None:
    with pytest.raises(ConfigError) as info:
        find_workbook(tmp_path)

    assert (info.value.file, info.value.path, info.value.found) == (tmp_path, "*.xlsx", "[]")
