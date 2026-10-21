"""Details from an uploaded file: only what Python can confirm is offered; the rest is unclear.
Offline: a fake reader stands in for Gemini."""

import io
import itertools

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.agent.agent import run_turn
from app.agent.document import (
    MAX_FILE_BYTES,
    Document,
    DocumentError,
    check_document,
    document_text,
    read_details,
    verified,
)
from app.agent.llm import Message, ProviderError, ProviderErrorKind, ToolCall
from app.agent.state import ConversationState
from app.agent.tools import ToolContext
from app.api.main import create_app
from app.api.sessions import SessionStore

PAYSLIP = """תלוש שכר - חודש 09/2026
שם העובד: ישראלה ישראלי    תפקיד: מנהלת מוצר
שכר בסיס: 18,500.00
עמלות: 0
הפרשת עובד לפנסיה: 6%
הפרשת עובד לקרן השתלמות: 2.5%
נקודות זיכוי: 2.75
רכב צמוד: קורולה קרוס    השתתפות עצמית: 300
"""


def entry(field: str, value: str, quote: str) -> dict:
    return {"field": field, "value": value, "quote": quote}


GOOD = [
    entry("occupation", "מנהלת מוצר", "תפקיד: מנהלת מוצר"),
    entry("base_salary", "18,500.00", "שכר בסיס: 18,500.00"),
    entry("commissions", "0", "עמלות: 0"),
    entry("pension_percent", "6%", "הפרשת עובד לפנסיה: 6%"),
    entry("training_fund_percent", "2.5%", "הפרשת עובד לקרן השתלמות: 2.5%"),
    entry("credit_points", "2.75", "נקודות זיכוי: 2.75"),
    entry("current_car", "קורולה קרוס", "רכב צמוד: קורולה קרוס"),
    entry("current_participation", "300", "השתתפות עצמית: 300"),
]


@pytest.fixture(scope="module")
def ctx() -> ToolContext:
    return ToolContext.load()


# --- What is certain ---------------------------------------------------------------------------------

def test_clear_values_are_offered_in_display_form(ctx: ToolContext) -> None:
    details = verified({"fields": GOOD, "unclear": ["bonus"]}, PAYSLIP, ctx)

    assert [(f.field, f.value) for f in details.found] == [
        ("occupation", "מנהל/ת מוצר"),
        ("base_salary", "18,500 ₪"),
        ("commissions", "0 ₪"),
        ("pension_percent", "6%"),
        ("training_fund_percent", "2.5%"),
        ("credit_points", "2.75"),
        ("current_car", "טויוטה יפן COROLLA CROSS (2025)"),
        ("current_participation", "300 ₪"),
    ]
    assert details.unclear == ("bonus",) and details.from_image is False
    assert details.found[1].label == "שכר בסיס חודשי (ברוטו)"


@pytest.mark.parametrize(
    "bad, why",
    [
        (entry("base_salary", "18,500", "שכר בסיס: 18,500 לחודש"), "the quote is not in the file"),
        (entry("base_salary", "19,000", "שכר בסיס: 18,500.00"), "the value is not in its quote"),
        (entry("base_salary", "בערך 18 אלף", "שכר בסיס: 18,500.00"), "not a number"),
        (entry("pension_percent", "600%", "הפרשת עובד לפנסיה: 6%"), "out of range / not in quote"),
        (entry("occupation", "מנהלת", "תפקיד: מנהלת מוצר"), "matches no single occupation"),
        (entry("current_car", "קורולה", "רכב צמוד: קורולה קרוס"), "several catalogue cars"),
        (entry("base_salary", "18,500", ""), "no quote"),
    ],
)
def test_anything_uncertain_is_unclear(ctx: ToolContext, bad: dict, why: str) -> None:
    details = verified({"fields": [bad], "unclear": []}, PAYSLIP, ctx)
    assert details.found == () and details.unclear == (bad["field"],), why


def test_two_values_for_one_field_are_unclear(ctx: ToolContext) -> None:
    text = PAYSLIP + "שכר בסיס מעודכן: 19,000\n"
    details = verified({"fields": [GOOD[1], entry("base_salary", "19,000", "שכר בסיס מעודכן: 19,000")], "unclear": []}, text, ctx)
    assert details.found == () and details.unclear == ("base_salary",)


def test_a_field_the_reader_also_calls_unclear_is_not_offered(ctx: ToolContext) -> None:
    details = verified({"fields": [GOOD[1]], "unclear": ["base_salary"]}, PAYSLIP, ctx)
    assert details.found == () and details.unclear == ("base_salary",)


def test_unknown_fields_and_malformed_replies_are_ignored(ctx: ToolContext) -> None:
    assert verified({"fields": [entry("id_number", "123", "ת.ז. 123")], "unclear": ["bank"]}, PAYSLIP, ctx).found == ()
    assert verified({"fields": "nonsense"}, PAYSLIP, ctx).found == ()
    assert verified([], PAYSLIP, ctx).found == ()  # type: ignore[arg-type]


def test_children_must_be_a_whole_number(ctx: ToolContext) -> None:
    text = "מספר ילדים: 2.5"
    details = verified({"fields": [entry("eligible_children", "2.5", "מספר ילדים: 2.5")], "unclear": []}, text, ctx)
    assert details.unclear == ("eligible_children",)


def test_from_an_image_the_quote_cannot_be_checked_but_everything_else_is(ctx: ToolContext) -> None:
    details = verified({"fields": [GOOD[1], entry("bonus", "5,000", "בונוס: 3,000")], "unclear": []}, None, ctx)
    assert [(f.field, f.value) for f in details.found] == [("base_salary", "18,500 ₪")]
    assert details.unclear == ("bonus",) and details.from_image is True


# --- Reading files ---------------------------------------------------------------------------------

def text_pdf(text: str) -> bytes:
    """A minimal one-page PDF with a text layer (ASCII), built by hand."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for number, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def xlsx(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def test_text_files_in_utf8_and_windows_hebrew() -> None:
    assert document_text(Document("a.txt", "שכר בסיס: 18,500".encode("utf-8-sig"))) == "שכר בסיס: 18,500"
    assert document_text(Document("a.txt", "שכר בסיס: 18,500".encode("cp1255"))) == "שכר בסיס: 18,500"


def test_csv_and_excel_become_lines_of_cells() -> None:
    assert document_text(Document("a.csv", "שדה,ערך\nשכר בסיס,18500\n".encode())) == "שדה | ערך\nשכר בסיס | 18500"
    assert document_text(Document("a.xlsx", xlsx([["שכר בסיס", 18500], ["נקודות זיכוי", 2.75]]))) == "שכר בסיס | 18500\nנקודות זיכוי | 2.75"


def test_a_pdf_with_text_is_read_and_a_scan_is_not() -> None:
    assert "Base salary 18500" in document_text(Document("p.pdf", text_pdf("Base salary 18500, credit points 2.75")))
    assert document_text(Document("scan.pdf", text_pdf(""))) is None  # no text layer: read as an image
    assert document_text(Document("photo.jpg", b"\xff\xd8\xff\xe0 jpeg")) is None


@pytest.mark.parametrize(
    "name, data, code",
    [
        ("a.txt", b"", "unreadable_file"),
        ("a.txt", b"x" * (MAX_FILE_BYTES + 1), "file_too_large"),
        ("a.exe", b"MZ", "unsupported_file"),
        ("a.png", b"not a png", "unsupported_file"),  # the extension alone is not trusted
        ("a.pdf", b"%PDF-1.4 broken", "unreadable_file"),
    ],
    ids=["empty", "too-large", "exe", "fake-png", "broken-pdf"],
)
def test_bad_files_are_refused(name: str, data: bytes, code: str) -> None:
    with pytest.raises(DocumentError) as raised:
        check_document(Document(name, data))
        document_text(Document(name, data))
    assert raised.value.code == code


class FakeReader:
    def __init__(self, reply: dict | Exception) -> None:
        self.reply, self.seen = reply, []

    def read(self, text, document):
        self.seen.append((text, document.name))
        if isinstance(self.reply, Exception):
            raise self.reply
        return self.reply


def test_the_reader_gets_the_text_or_else_the_file(ctx: ToolContext) -> None:
    reader = FakeReader({"fields": [], "unclear": []})
    read_details(Document("a.txt", PAYSLIP.encode()), reader, ctx)
    read_details(Document("p.png", b"\x89PNG..."), reader, ctx)
    assert reader.seen == [(PAYSLIP, "a.txt"), (None, "p.png")]


# --- The API ---------------------------------------------------------------------------------------

_ids = itertools.count()


class Scripted:
    name = "scripted"

    def __init__(self, *steps: Message) -> None:
        self._steps = list(steps)

    def complete(self, system, messages, tools) -> Message:
        assert self._steps, "the agent asked the provider for more replies than scripted"
        return self._steps.pop(0)


def app_with(ctx: ToolContext, reader, provider=None) -> TestClient:
    return TestClient(create_app(provider or Scripted(), ctx, SessionStore(), signing_key=b"k" * 32, document_reader=reader))


def upload(client: TestClient, session_id: str, name: str, data: bytes):
    return client.post(f"/api/sessions/{session_id}/document", files={"file": (name, data)})


def test_upload_returns_only_what_is_certain_and_saves_nothing(ctx: ToolContext) -> None:
    store = SessionStore()
    client = TestClient(create_app(Scripted(), ctx, store, signing_key=b"k" * 32,
                                   document_reader=FakeReader({"fields": GOOD[:2], "unclear": ["bonus"]})))
    session_id = client.post("/api/sessions").json()["session_id"]

    response = upload(client, session_id, "payslip.txt", PAYSLIP.encode())

    assert response.status_code == 200 and response.headers["cache-control"] == "no-store"
    assert response.json() == {
        "found": [
            {"field": "occupation", "label": "תפקיד", "value": "מנהל/ת מוצר"},
            {"field": "base_salary", "label": "שכר בסיס חודשי (ברוטו)", "value": "18,500 ₪"},
        ],
        "unclear": [{"field": "bonus", "label": "בונוס חודשי"}],
        "from_image": False,
    }
    assert store.get(session_id).state == ConversationState(session_id)  # nothing saved before confirmation


@pytest.mark.parametrize(
    "name, data, status, code",
    [
        ("big.txt", b"x" * (MAX_FILE_BYTES + 10), 413, "file_too_large"),
        ("virus.exe", b"MZ", 415, "unsupported_file"),
        ("empty.txt", b"", 422, "unreadable_file"),
    ],
    ids=["too-large", "exe", "empty"],
)
def test_upload_errors(ctx: ToolContext, name: str, data: bytes, status: int, code: str) -> None:
    client = app_with(ctx, FakeReader({"fields": [], "unclear": []}))
    session_id = client.post("/api/sessions").json()["session_id"]
    response = upload(client, session_id, name, data)
    assert (response.status_code, response.json()["error"]["code"]) == (status, code)


def test_upload_needs_a_session_and_a_working_reader(ctx: ToolContext) -> None:
    down = ProviderError("gemini", ProviderErrorKind.RATE_LIMITED, "429")
    client = app_with(ctx, FakeReader(down))
    assert upload(client, "0" * 32, "a.txt", b"x").status_code == 404
    session_id = client.post("/api/sessions").json()["session_id"]
    response = upload(client, session_id, "a.txt", PAYSLIP.encode())
    assert (response.status_code, response.json()["error"]["code"]) == (503, "document_unavailable")
    assert "429" not in response.text


def test_the_confirmation_in_the_chat_is_grounded_and_saved(ctx: ToolContext) -> None:
    # The frontend sends the confirmed details as the employee's own words; the advisor saves them with
    # the usual tools, and the grounding checks accept them because the employee wrote them.
    details = verified({"fields": GOOD, "unclear": ["bonus"]}, PAYSLIP, ctx)
    message = "\n".join(
        ["אלה הפרטים מהקובץ שהעליתי, ובדקתי שהם נכונים:", *(f"• {f.label}: {f.value}" for f in details.found),
         "לא היו ברורים בקובץ: בונוס חודשי."]
    )
    save = ToolCall("d-1", "save_employee_details", {
        "occupation": "מנהל/ת מוצר", "base_salary": 18500, "commissions": 0, "pension_percent": 6,
        "training_fund_percent": 2.5, "credit_points": 2.75,
    })
    current = ToolCall("d-2", "set_current_car", {"car": "טויוטה יפן COROLLA CROSS (2025)", "participation": 300})
    provider = Scripted(Message.assistant(tool_calls=[save, current]), Message.assistant("מה הבונוס החודשי שלך?"))

    result = run_turn(ConversationState("c-1"), message, provider, ctx)

    saved = result.state.details
    assert (saved.occupation, saved.base_salary, saved.pension_rate, saved.credit_points) == ("מנהל/ת מוצר", 18500, 0.06, 2.75)
    assert saved.current is not None and saved.current.car_id == "CAR-09" and saved.current.participation == 300
    assert saved.bonus is None  # unclear in the file: asked again
